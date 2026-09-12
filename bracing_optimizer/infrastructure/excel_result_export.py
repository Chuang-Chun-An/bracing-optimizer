"""Write visible Solver material results to a formatted Excel workbook."""

from __future__ import annotations

import math
import os
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from bracing_optimizer.application.project_results import MaterialDetailRow


DETAIL_SHEET_NAME = "材料明細"
SUMMARY_SHEET_NAME = "材料彙總"
DETAIL_HEADERS = (
    "序號",
    "材料用途",
    "所屬構件編號",
    "成果方案",
    "分區",
    "構件內段次",
    "材料類型",
    "材料規格",
    "長度(mm)",
    "數量",
)
SUMMARY_HEADERS = (
    "材料用途",
    "材料規格",
    "長度(mm)",
    "使用數量",
    "庫存數量",
    "剩餘數量",
)

TITLE_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
LABEL_FILL = PatternFill("solid", fgColor="D9EAF7")
SHORTAGE_FILL = PatternFill("solid", fgColor="FCE8E6")
WHITE_FONT = Font(color="FFFFFF", bold=True, name="Microsoft JhengHei")
TITLE_FONT = Font(color="FFFFFF", bold=True, size=16, name="Microsoft JhengHei")
BODY_FONT = Font(color="1F2937", size=10, name="Microsoft JhengHei")
LABEL_FONT = Font(color="17365D", bold=True, name="Microsoft JhengHei")
SHORTAGE_FONT = Font(color="B91C1C", bold=True, name="Microsoft JhengHei")
LIGHT_SIDE = Side(style="thin", color="CBD5E1")
SECTION_BORDER = Border(bottom=Side(style="medium", color="9FBAD0"))

MATERIAL_TYPE_LABELS = {
    "steel": "鋼材",
    "shim": "調整塊",
    "jack": "千斤頂",
}


class ExcelResultExportError(ValueError):
    """Raised when a result workbook cannot be safely created."""


@dataclass(frozen=True)
class ExcelResultExportReport:
    output_path: str
    detail_row_count: int
    material_quantity: int | float
    summary_row_count: int


def _finite_number(value, *, field_name: str, nonnegative: bool = False):
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ExcelResultExportError(f"{field_name}不是有效數字：{value!r}") from exc
    if not math.isfinite(number) or (number < 0 if nonnegative else number <= 0):
        comparator = "不可小於 0" if nonnegative else "必須大於 0"
        raise ExcelResultExportError(f"{field_name}{comparator}：{value!r}")
    return int(number) if number.is_integer() else number


def _signed_finite_number(value, *, field_name: str):
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ExcelResultExportError(f"{field_name}不是有效數字：{value!r}") from exc
    if not math.isfinite(number):
        raise ExcelResultExportError(f"{field_name}不是有效數字：{value!r}")
    return int(number) if number.is_integer() else number


def _display_material_spec(value) -> str:
    return str(value or "").strip() or "未設定"


def _metadata_row(
    worksheet,
    *,
    project_name: str,
    exported_at: datetime,
    total_formula: str | None = None,
) -> None:
    worksheet["A2"] = "專案"
    worksheet["B2"] = project_name
    worksheet.merge_cells("B2:C2")
    worksheet["D2"] = "匯出時間"
    worksheet["E2"] = exported_at
    worksheet.merge_cells("E2:G2" if total_formula is not None else "E2:F2")
    worksheet["E2"].number_format = "yyyy-mm-dd hh:mm:ss"
    if total_formula is not None:
        worksheet["H2"] = "材料總數"
        worksheet["I2"] = total_formula
        worksheet.merge_cells("I2:J2")
        worksheet["I2"].number_format = "#,##0.##"

    for coordinate in ("A2", "D2", "H2"):
        cell = worksheet[coordinate]
        if cell.value is None:
            continue
        cell.fill = LABEL_FILL
        cell.font = LABEL_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for coordinate in ("B2", "E2", "I2"):
        cell = worksheet[coordinate]
        if cell.value is None:
            continue
        cell.font = BODY_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[2].height = 24


def _style_sheet_title(worksheet, title: str, last_column: int) -> None:
    end_column = get_column_letter(last_column)
    worksheet.merge_cells(f"A1:{end_column}1")
    title_cell = worksheet["A1"]
    title_cell.value = title
    title_cell.fill = TITLE_FILL
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    title_cell.border = SECTION_BORDER
    worksheet.row_dimensions[1].height = 30
    worksheet.sheet_view.showGridLines = False


def _style_header(worksheet, *, header_row: int, column_count: int) -> None:
    for cell in worksheet[header_row][:column_count]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=LIGHT_SIDE)
    worksheet.row_dimensions[header_row].height = 24


def _add_table(worksheet, *, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _configure_printing(worksheet, *, column_count: int) -> None:
    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = (
        f"A4:{get_column_letter(column_count)}{worksheet.max_row}"
    )
    worksheet.print_title_rows = "1:4"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.outlinePr.summaryBelow = True


def _write_detail_sheet(
    workbook: Workbook,
    rows: Sequence[MaterialDetailRow],
    *,
    project_name: str,
    exported_at: datetime,
) -> tuple[int | float, int]:
    worksheet = workbook.active
    worksheet.title = DETAIL_SHEET_NAME
    _style_sheet_title(
        worksheet,
        f"{project_name}－配置材料明細",
        len(DETAIL_HEADERS),
    )
    first_data_row = 5
    last_data_row = first_data_row + len(rows) - 1
    _metadata_row(
        worksheet,
        project_name=project_name,
        exported_at=exported_at,
        total_formula=f"=SUM(J{first_data_row}:J{last_data_row})",
    )
    for column, header in enumerate(DETAIL_HEADERS, start=1):
        worksheet.cell(4, column, header)

    material_quantity = 0
    for sequence, row in enumerate(rows, start=1):
        length = _finite_number(
            row.length,
            field_name=f"{row.usage} {row.member_id} 長度",
        )
        quantity = _finite_number(
            row.quantity,
            field_name=f"{row.usage} {row.member_id} 數量",
        )
        material_quantity += quantity
        worksheet.append([
            sequence,
            row.usage,
            row.member_id,
            row.result_id,
            row.zoning or "—",
            row.piece_index,
            MATERIAL_TYPE_LABELS.get(row.material_type, row.material_type or "其他"),
            _display_material_spec(row.material_spec),
            length,
            quantity,
        ])

    _style_header(
        worksheet,
        header_row=4,
        column_count=len(DETAIL_HEADERS),
    )
    for data_row in worksheet.iter_rows(
        min_row=first_data_row,
        max_row=last_data_row,
        min_col=1,
        max_col=len(DETAIL_HEADERS),
    ):
        for cell in data_row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")
        for column in (1, 2, 5, 6, 7, 10):
            data_row[column - 1].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
        data_row[8].alignment = Alignment(horizontal="right", vertical="center")
        data_row[9].alignment = Alignment(horizontal="right", vertical="center")
        data_row[8].number_format = "#,##0.##"
        data_row[9].number_format = "#,##0.##"

    widths = (8, 12, 18, 22, 14, 12, 14, 22, 14, 10)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    _add_table(
        worksheet,
        name="MaterialDetailsTable",
        reference=f"A4:J{last_data_row}",
    )
    _configure_printing(worksheet, column_count=len(DETAIL_HEADERS))
    return material_quantity, last_data_row


def _write_summary_sheet(
    workbook: Workbook,
    rows: Sequence[Mapping],
    *,
    project_name: str,
    exported_at: datetime,
) -> None:
    worksheet = workbook.create_sheet(SUMMARY_SHEET_NAME)
    _style_sheet_title(
        worksheet,
        f"{project_name}－材料用量與庫存彙總",
        len(SUMMARY_HEADERS),
    )
    _metadata_row(
        worksheet,
        project_name=project_name,
        exported_at=exported_at,
    )
    for column, header in enumerate(SUMMARY_HEADERS, start=1):
        worksheet.cell(4, column, header)

    first_data_row = 5
    for row in rows:
        length = _finite_number(row.get("length"), field_name="彙總長度")
        used_quantity = _finite_number(
            row.get("used_qty"),
            field_name="使用數量",
            nonnegative=True,
        )
        inventory_quantity = _finite_number(
            row.get("inventory_qty"),
            field_name="庫存數量",
            nonnegative=True,
        )
        remaining_quantity = _signed_finite_number(
            row.get("remaining_qty"),
            field_name="剩餘數量",
        )
        worksheet.append([
            str(row.get("usage", "") or "").strip(),
            _display_material_spec(row.get("material_spec")),
            length,
            used_quantity,
            inventory_quantity,
            remaining_quantity,
        ])

    last_data_row = worksheet.max_row
    _style_header(
        worksheet,
        header_row=4,
        column_count=len(SUMMARY_HEADERS),
    )
    if rows:
        for data_row in worksheet.iter_rows(
            min_row=first_data_row,
            max_row=last_data_row,
            min_col=1,
            max_col=len(SUMMARY_HEADERS),
        ):
            for cell in data_row:
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="center")
            data_row[0].alignment = Alignment(horizontal="center", vertical="center")
            for cell in data_row[2:]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.##"
        _add_table(
            worksheet,
            name="MaterialSummaryTable",
            reference=f"A4:F{last_data_row}",
        )
        worksheet.conditional_formatting.add(
            f"F{first_data_row}:F{last_data_row}",
            CellIsRule(
                operator="lessThan",
                formula=["0"],
                fill=SHORTAGE_FILL,
                font=SHORTAGE_FONT,
            ),
        )
    widths = (14, 24, 14, 14, 14, 14)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    _configure_printing(worksheet, column_count=len(SUMMARY_HEADERS))


def _validate_staged_workbook(
    staged_path: Path,
    *,
    expected_detail_rows: int,
    expected_summary_rows: int,
) -> None:
    try:
        workbook = load_workbook(staged_path, read_only=True, data_only=False)
        try:
            if workbook.sheetnames != [DETAIL_SHEET_NAME, SUMMARY_SHEET_NAME]:
                raise ExcelResultExportError("Excel 工作表結構不完整。")
            detail_sheet = workbook[DETAIL_SHEET_NAME]
            summary_sheet = workbook[SUMMARY_SHEET_NAME]
            detail_headers = tuple(
                detail_sheet.cell(4, column).value
                for column in range(1, len(DETAIL_HEADERS) + 1)
            )
            summary_headers = tuple(
                summary_sheet.cell(4, column).value
                for column in range(1, len(SUMMARY_HEADERS) + 1)
            )
            if detail_headers != DETAIL_HEADERS or summary_headers != SUMMARY_HEADERS:
                raise ExcelResultExportError("Excel 欄位結構驗證失敗。")
            if detail_sheet.max_row != expected_detail_rows + 4:
                raise ExcelResultExportError("Excel 材料明細筆數驗證失敗。")
            if summary_sheet.max_row != max(4, expected_summary_rows + 4):
                raise ExcelResultExportError("Excel 材料彙總筆數驗證失敗。")
        finally:
            workbook.close()
    except ExcelResultExportError:
        raise
    except Exception as exc:
        raise ExcelResultExportError(f"Excel 暫存檔驗證失敗：{exc}") from exc


class ExcelResultExporter:
    """Create and atomically deliver one Excel material result workbook."""

    def export(
        self,
        output_path: str | Path,
        detail_rows: Sequence[MaterialDetailRow],
        summary_rows: Sequence[Mapping],
        *,
        project_name: str,
        exported_at: datetime | None = None,
    ) -> ExcelResultExportReport:
        details = tuple(detail_rows)
        summaries = tuple(summary_rows)
        if not details:
            raise ExcelResultExportError("目前沒有可匯出的材料明細。")

        destination = Path(output_path)
        if destination.suffix.lower() != ".xlsx":
            raise ExcelResultExportError("Excel 成果檔必須使用 .xlsx 副檔名。")
        if not destination.parent.is_dir():
            raise ExcelResultExportError(
                f"輸出資料夾不存在：{destination.parent}"
            )

        exported_at = exported_at or datetime.now()
        display_name = str(project_name or "").strip() or "未命名專案"
        workbook = Workbook()
        workbook.properties.title = f"{display_name}－配置材料明細"
        workbook.properties.subject = "目前顯示配置成果的逐根材料明細與庫存彙總"
        workbook.properties.creator = "Support Distribution UV"
        material_quantity, _ = _write_detail_sheet(
            workbook,
            details,
            project_name=display_name,
            exported_at=exported_at,
        )
        _write_summary_sheet(
            workbook,
            summaries,
            project_name=display_name,
            exported_at=exported_at,
        )

        staged_path = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="wb",
                prefix=f".{destination.stem}-",
                suffix=".xlsx",
                dir=destination.parent,
                delete=False,
            ) as staged_file:
                staged_path = Path(staged_file.name)
            workbook.save(staged_path)
            workbook.close()
            _validate_staged_workbook(
                staged_path,
                expected_detail_rows=len(details),
                expected_summary_rows=len(summaries),
            )
            os.replace(staged_path, destination)
            staged_path = None
        except ExcelResultExportError:
            raise
        except (OSError, ValueError) as exc:
            raise ExcelResultExportError(f"無法建立 Excel 成果檔：{exc}") from exc
        finally:
            workbook.close()
            if staged_path is not None:
                try:
                    staged_path.unlink(missing_ok=True)
                except OSError:
                    pass

        return ExcelResultExportReport(
            output_path=str(destination),
            detail_row_count=len(details),
            material_quantity=material_quantity,
            summary_row_count=len(summaries),
        )


__all__ = [
    "DETAIL_HEADERS",
    "DETAIL_SHEET_NAME",
    "ExcelResultExportError",
    "ExcelResultExportReport",
    "ExcelResultExporter",
    "SUMMARY_HEADERS",
    "SUMMARY_SHEET_NAME",
]
