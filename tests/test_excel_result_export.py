import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from bracing_optimizer.application.project_results import MaterialDetailRow
from bracing_optimizer.infrastructure.excel_result_export import (
    DETAIL_HEADERS,
    ExcelResultExportError,
    ExcelResultExporter,
    SUMMARY_HEADERS,
)


class ExcelResultExporterTests(unittest.TestCase):
    @staticmethod
    def detail_rows():
        return [
            MaterialDetailRow(
                result_id="W1-plan-1",
                usage="圍令",
                member_id="W1",
                zoning="",
                piece_index=1,
                material_type="steel",
                material_spec="H350",
                length=9000,
            ),
            MaterialDetailRow(
                result_id="Z1",
                usage="支撐",
                member_id="S1",
                zoning="Z1",
                piece_index=2,
                material_type="jack",
                material_spec="",
                length=600,
            ),
        ]

    @staticmethod
    def summary_rows():
        return [
            {
                "usage": "圍令",
                "material_spec": "H350",
                "length": 9000,
                "used_qty": 2,
                "inventory_qty": 1,
                "remaining_qty": -1,
            },
        ]

    def test_export_writes_typed_detail_and_summary_tables(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "配置材料明細.xlsx"

            report = ExcelResultExporter().export(
                output_path,
                self.detail_rows(),
                self.summary_rows(),
                project_name="測試專案",
                exported_at=datetime(2026, 9, 11, 10, 30),
            )

            workbook = load_workbook(output_path, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["材料明細", "材料彙總"])
                detail = workbook["材料明細"]
                summary = workbook["材料彙總"]
                self.assertEqual(
                    tuple(detail.cell(4, column).value for column in range(1, 11)),
                    DETAIL_HEADERS,
                )
                self.assertEqual(
                    tuple(summary.cell(4, column).value for column in range(1, 7)),
                    SUMMARY_HEADERS,
                )
                self.assertEqual(
                    [detail.cell(5, column).value for column in range(1, 11)],
                    [
                        1,
                        "圍令",
                        "W1",
                        "W1-plan-1",
                        "—",
                        1,
                        "鋼材",
                        "H350",
                        9000,
                        1,
                    ],
                )
                self.assertEqual(detail["G6"].value, "千斤頂")
                self.assertEqual(detail["H6"].value, "未設定")
                self.assertEqual(detail["I2"].value, "=SUM(J5:J6)")
                self.assertEqual(detail.freeze_panes, "A5")
                self.assertIn("MaterialDetailsTable", detail.tables)
                self.assertEqual(detail["I5"].number_format, "#,##0.##")
                self.assertEqual(summary["F5"].value, -1)
                self.assertIn("MaterialSummaryTable", summary.tables)
                self.assertEqual(summary.freeze_panes, "A5")
                self.assertFalse(detail.sheet_view.showGridLines)
                self.assertFalse(summary.sheet_view.showGridLines)
            finally:
                workbook.close()

        self.assertEqual(report.output_path, str(output_path))
        self.assertEqual(report.detail_row_count, 2)
        self.assertEqual(report.material_quantity, 2)
        self.assertEqual(report.summary_row_count, 1)

    def test_invalid_material_does_not_replace_existing_destination(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "existing.xlsx"
            output_path.write_bytes(b"existing workbook placeholder")
            invalid_row = MaterialDetailRow(
                result_id="Z1",
                usage="支撐",
                member_id="S1",
                zoning="Z1",
                piece_index=1,
                material_type="steel",
                material_spec="H400",
                length=0,
            )

            with self.assertRaises(ExcelResultExportError):
                ExcelResultExporter().export(
                    output_path,
                    [invalid_row],
                    [],
                    project_name="測試專案",
                )

            self.assertEqual(
                output_path.read_bytes(),
                b"existing workbook placeholder",
            )

    def test_export_requires_material_rows_and_xlsx_extension(self):
        exporter = ExcelResultExporter()
        with tempfile.TemporaryDirectory() as temp_dir:
            with self.assertRaisesRegex(ExcelResultExportError, "沒有可匯出"):
                exporter.export(
                    Path(temp_dir) / "empty.xlsx",
                    [],
                    [],
                    project_name="測試專案",
                )
            with self.assertRaisesRegex(ExcelResultExportError, "xlsx"):
                exporter.export(
                    Path(temp_dir) / "materials.xls",
                    self.detail_rows(),
                    [],
                    project_name="測試專案",
                )


if __name__ == "__main__":
    unittest.main()
